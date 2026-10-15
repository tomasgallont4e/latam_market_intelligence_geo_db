"""Publish aggregate, national TUNI admission outcomes without institution records."""
from __future__ import annotations

import json
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl


if len(sys.argv) != 4:
    raise SystemExit('Usage: python scripts/build-peru-tuni-admissions.py <Ingresantes.xlsx> <output-directory> <map-fact.json>')

source, outdir, mapfact = map(Path, sys.argv[1:])
outdir.mkdir(parents=True, exist_ok=True)


def canonical(value: object) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '').strip()) if unicodedata.category(c) != 'Mn').upper()


def number(value: object) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = book[book.sheetnames[0]]
rows = sheet.iter_rows(values_only=True)
next(rows, None)
next(rows, None)
header = next(rows, None)
if not header or canonical(header[0]) != 'NRO' or canonical(header[-1]) != 'TOTAL DE INGRESANTES':
    raise SystemExit('Expected the public TUNI Ingresantes export')

by_group: dict[str, int] = defaultdict(int)
total = ict_total = accepted_rows = source_rows = 0
for row in rows:
    source_rows += 1
    if not row or len(row) < 7 or not isinstance(row[0], (int, float)):
        continue
    value = number(row[6])
    group = canonical(row[4])
    total += value
    by_group[group or 'SIN_CLASIFICACION'] += value
    if group in {'SISTEMAS Y COMPUTO', 'INFORMATICA'}:
        ict_total += value
    accepted_rows += 1

metadata = {
    'country': 'PE',
    'source': 'SUNEDU TUNI, Comunidad estudiantil / Ingresantes',
    'source_url': 'https://www.tuni.pe/comunidad/estudiantes#ingresantes',
    'downloaded_at': '2026-10-15',
    'reference_period': '2025',
    'source_rows': source_rows,
    'accepted_rows': accepted_rows,
    'caveat': 'TUNI public export is an aggregate by institution and programme group. It has no local, district or department field, so these values are national only and must not be allocated geographically. ICT is the TUNI programme groups Sistemas y Cómputo or Informática only.',
}
aggregate = {'metric': 'entrants', 'total_entrants': total, 'ict_entrants': ict_total, 'programme_group_totals': [{'programme_group': key, 'entrants': value} for key, value in sorted(by_group.items())], **metadata}
facts = [
    {'layer': 'C', 'metric_id': 'C_TUNI_INGRESS_TOTAL_2025', 'scope_type': 'country', 'scope_id': 'PE', 'group_id': '', 'persona_id': '', 'indicator': 'Ingresantes reportados (TUNI/SUNEDU)', 'category': 'higher_education', 'value': total, 'unit': 'students', 'period': '2025', 'comparison_period': '', 'numerator': None, 'denominator': None, 'status': 'published_aggregate_national', 'source_id': 'SUNEDU_TUNI_INGRESANTES_2025', 'method': 'sum of public TUNI programme rows', 'limitations': metadata['caveat']},
    {'layer': 'C', 'metric_id': 'C_TUNI_INGRESS_ICT_2025', 'scope_type': 'country', 'scope_id': 'PE', 'group_id': '', 'persona_id': '', 'indicator': 'Ingresantes TIC reportados (TUNI/SUNEDU)', 'category': 'higher_education', 'value': ict_total, 'unit': 'students', 'period': '2025', 'comparison_period': '', 'numerator': None, 'denominator': None, 'status': 'published_aggregate_national', 'source_id': 'SUNEDU_TUNI_INGRESANTES_2025', 'method': 'sum of public TUNI rows classified Sistemas y Cómputo or Informática', 'limitations': metadata['caveat']},
]
(outdir / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf8')
(outdir / 'ingress_national_2025.json').write_text(json.dumps(aggregate, ensure_ascii=False, indent=2), encoding='utf8')
mapfact.write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps({'total_entrants': total, 'ict_entrants': ict_total, **metadata}, ensure_ascii=False))
