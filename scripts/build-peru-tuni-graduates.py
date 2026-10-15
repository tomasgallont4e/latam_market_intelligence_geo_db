"""Aggregate the public SUNEDU TUNI graduate export without retaining local addresses."""
from __future__ import annotations

import json
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

if len(sys.argv) != 4:
    raise SystemExit('Usage: python scripts/build-peru-tuni-graduates.py <Egresados.xlsx> <output-directory> <map-directory>')

source, outdir, mapdir = map(Path, sys.argv[1:])
outdir.mkdir(parents=True, exist_ok=True)
mapdir.mkdir(parents=True, exist_ok=True)
departments = ['AMAZONAS','ANCASH','APURIMAC','AREQUIPA','AYACUCHO','CAJAMARCA','CALLAO','CUSCO','HUANCAVELICA','HUANUCO','ICA','JUNIN','LA LIBERTAD','LAMBAYEQUE','LIMA','LORETO','MADRE DE DIOS','MOQUEGUA','PASCO','PIURA','PUNO','SAN MARTIN','TACNA','TUMBES','UCAYALI']

def canonical(value: object) -> str:
    return ''.join(c for c in unicodedata.normalize('NFD', str(value or '').strip()) if unicodedata.category(c) != 'Mn').upper()

def number(value: object) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0

def blank() -> dict:
    return {'graduates': 0, 'ict_graduates': 0, 'records': 0, 'ict_records': 0, 'institutions': set(), 'locals': set()}

def row_for(key: tuple[str, ...], value: dict, geography: tuple[str, ...]) -> dict:
    return {**dict(zip(geography, key)), 'graduates': value['graduates'], 'ict_graduates': value['ict_graduates'], 'programme_records': value['records'], 'ict_programme_records': value['ict_records'], 'institutions_reporting': len(value['institutions']), 'locals_reporting': len(value['locals']), 'reference_period': '2024', 'source_id': 'SUNEDU_TUNI_GRADUATES_2024', 'scope': 'reported_higher_education_graduates'}

book = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = book[book.sheetnames[0]]
rows = sheet.iter_rows(values_only=True)
title = next(rows, None)
next(rows, None)
header = next(rows, None)
if not title or canonical(title[0]) != 'EGRESADOS' or not header or canonical(header[0]) != 'NRO' or canonical(header[-1]) not in {'TOTAL DE EGRESADOS', 'TOTAL DE ESTUDIANTES'}:
    raise SystemExit('Expected the public TUNI Egresados export')

by_department: dict[str, dict] = defaultdict(blank)
by_district: dict[tuple[str, str, str], dict] = defaultdict(blank)
source_rows = accepted_rows = 0
for values in rows:
    source_rows += 1
    if not values or len(values) < 12 or not isinstance(values[0], (int, float)):
        continue
    department, province, district = canonical(values[7]), canonical(values[8]), canonical(values[9])
    if str(values[3] or '').strip() != '2024' or department not in departments or not province or not district:
        continue
    graduates = number(values[11])
    group = canonical(values[4])
    ict = group in {'SISTEMAS Y COMPUTO', 'INFORMATICA'}
    for target in (by_department[department], by_district[(department, province, district)]):
        target['graduates'] += graduates; target['records'] += 1
        target['institutions'].add(str(values[1] or '').strip()); target['locals'].add(str(values[10] or '').strip())
        if ict:
            target['ict_graduates'] += graduates; target['ict_records'] += 1
    accepted_rows += 1

department_rows = [row_for((name,), by_department[name], ('department',)) for name in departments]
district_rows = [row_for(key, value, ('department', 'province', 'district')) for key, value in sorted(by_district.items())]
metadata = {'country': 'PE', 'source': 'SUNEDU TUNI, Comunidad estudiantil / Egresados', 'source_url': 'https://www.tuni.pe/comunidad/estudiantes#egresados', 'downloaded_at': '2026-10-15', 'reference_period': '2024', 'source_rows': source_rows, 'accepted_rows': accepted_rows, 'departments': len(department_rows), 'districts': len(district_rows), 'caveat': 'TUNI export aggregates reported graduates by programme and local. It contains no individual people. It covers institutions reporting to SIU and is not a census of all Peruvian tertiary education. ICT is the TUNI programme groups Sistemas y Cómputo or Informática only.'}
(outdir / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf8')
(outdir / 'tuni_graduates_by_department_2024.json').write_text(json.dumps(department_rows, ensure_ascii=False, indent=2), encoding='utf8')
(outdir / 'tuni_graduates_by_district_2024.json').write_text(json.dumps(district_rows, ensure_ascii=False, indent=2), encoding='utf8')
(mapdir / 'fact_tuni_graduates_by_department_2024.json').write_text(json.dumps(department_rows, ensure_ascii=False, indent=2), encoding='utf8')
(mapdir / 'fact_tuni_graduates_by_district_2024.json').write_text(json.dumps(district_rows, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps(metadata, ensure_ascii=False))
