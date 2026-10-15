"""Aggregate the public SUNEDU TUNI enrollment export without retaining local addresses."""
from __future__ import annotations

import json
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl


if len(sys.argv) != 3:
    raise SystemExit('Usage: python scripts/build-peru-tuni-enrollment.py <Matriculados.xlsx> <output-directory>')

source, outdir = map(Path, sys.argv[1:3])
outdir.mkdir(parents=True, exist_ok=True)
departments = ['AMAZONAS','ANCASH','APURIMAC','AREQUIPA','AYACUCHO','CAJAMARCA','CALLAO','CUSCO','HUANCAVELICA','HUANUCO','ICA','JUNIN','LA LIBERTAD','LAMBAYEQUE','LIMA','LORETO','MADRE DE DIOS','MOQUEGUA','PASCO','PIURA','PUNO','SAN MARTIN','TACNA','TUMBES','UCAYALI']


def canonical(value: object) -> str:
    text = str(value or '').strip()
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn').upper()


def number(value: object) -> int:
    try:
        return int(float(value or 0))
    except (TypeError, ValueError):
        return 0


def empty() -> dict:
    return {'students': 0, 'ict_students': 0, 'records': 0, 'ict_records': 0, 'institutions': set(), 'locals': set()}


def serialise(key: tuple[str, ...], value: dict, geography: tuple[str, ...]) -> dict:
    return {
        **dict(zip(geography, key)),
        'students': value['students'],
        'ict_students': value['ict_students'],
        'programme_records': value['records'],
        'ict_programme_records': value['ict_records'],
        'institutions_reporting': len(value['institutions']),
        'locals_reporting': len(value['locals']),
        'reference_period': '2025-2',
        'source_id': 'SUNEDU_TUNI_ENROLLMENT_2025_2',
        'scope': 'reported_higher_education_enrollment',
    }


workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = workbook[workbook.sheetnames[0]]
rows = sheet.iter_rows(values_only=True)
# The export has a title and timestamp, followed by this published header.
next(rows, None)
next(rows, None)
header = next(rows, None)
if not header or canonical(header[0]) != 'NRO' or canonical(header[3]) != 'PERIODO':
    raise SystemExit('Unexpected TUNI export format')

by_department: dict[str, dict] = defaultdict(empty)
by_district: dict[tuple[str, str, str], dict] = defaultdict(empty)
source_rows = accepted_rows = 0
for row in rows:
    source_rows += 1
    if not row or not isinstance(row[0], (int, float)):
        continue
    # Nro, institution, management, period, programme group, programme, level,
    # local department/province/district/local, total students.
    period = str(row[3] or '').strip()
    department, province, district = (canonical(row[7]), canonical(row[8]), canonical(row[9]))
    if period != '2025-2' or department not in departments or not province or not district:
        continue
    students = number(row[11])
    group = canonical(row[4])
    # TUNI's published programme-group taxonomy identifies computing explicitly.
    is_ict = group in {'SISTEMAS Y COMPUTO', 'INFORMATICA'}
    institution = str(row[1] or '').strip()
    local = str(row[10] or '').strip()
    for target in (by_department[department], by_district[(department, province, district)]):
        target['students'] += students
        target['records'] += 1
        target['institutions'].add(institution)
        target['locals'].add(local)
        if is_ict:
            target['ict_students'] += students
            target['ict_records'] += 1
    accepted_rows += 1

department_rows = [serialise((department,), by_department[department], ('department',)) for department in departments]
district_rows = [serialise(key, value, ('department', 'province', 'district')) for key, value in sorted(by_district.items())]
metadata = {
    'country': 'PE',
    'source': 'SUNEDU TUNI, Comunidad estudiantil / Matriculados',
    'source_url': 'https://www.tuni.pe/comunidad/estudiantes',
    'downloaded_at': '2026-10-15',
    'reference_period': '2025-2',
    'source_rows': source_rows,
    'accepted_rows': accepted_rows,
    'departments': len(department_rows),
    'districts': len(district_rows),
    'caveat': 'TUNI export aggregates reported enrollment by programme and local. It contains no individual students. It covers institutions reporting to SIU and is not a census of all Peruvian tertiary education. ICT is the TUNI programme groups Sistemas y Cómputo or Informática only.',
}
(outdir / 'metadata.json').write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding='utf8')
(outdir / 'tuni_enrollment_by_department_2025_2.json').write_text(json.dumps(department_rows, ensure_ascii=False, indent=2), encoding='utf8')
(outdir / 'tuni_enrollment_by_district_2025_2.json').write_text(json.dumps(district_rows, ensure_ascii=False, indent=2), encoding='utf8')
print(json.dumps(metadata, ensure_ascii=False))
